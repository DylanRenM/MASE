#!/usr/bin/env python3
"""Verify the editable MASE v2.4 training deck against approved V1 layouts."""

from __future__ import annotations

import argparse
import json
from hashlib import sha256
from pathlib import Path
import re
from typing import Any

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "training/mase-framework/mase-training-v2.4.yaml"
DEFAULT_DECK = ROOT / "training/mase-framework/MASE框架培训讲义V2.4.pptx"
V1 = ROOT / "training/mase-framework/MASE框架培训讲义V1.pptx"
V1_SHA256 = "fcdf4bd0b8fe3adb5316c8dbccf6e70d91b0274857a5f97827709053886fe188"

REQUIRED_TOPICS = (
    "MASE v2.4",
    "风险自适应",
    "Lite",
    "Standard",
    "Strict",
    "六步开发主线",
    "三种过程档位",
    "轻量档（Lite）",
    "标准档（Standard）",
    "严格档（Strict）",
    "Capability",
    "门禁运行器",
    "候选冻结",
    "mase-state.yaml",
    ".mase/gates.yaml",
    "Brownfield",
    "隔离环境",
    "上下文路由",
    "属性测试",
    "最小反例",
    "可选发布附加流程",
    "打包、发布、部署、线上验证或恢复",
    "不可变制品",
    "target_ready",
    "live_verified",
    "release-software",
    "观察与恢复",
    "framework_contract",
    "产品独立仓库",
    "已安装 CLI",
    "版本化 Schema",
    "临时采用项目",
    "六项开发原则",
    "验证职责分层",
    "原则落地：三次回看",
    "Spec 与测试设计",
    "独立测试用例用 Test ID 追踪",
    "风险与契约驱动设计",
    "证据新鲜度闭环",
    "需求到证据的追踪链",
    "候选 → 制品 → 目标 → 线上 → 观察",
    "门禁计划",
    "固定代码、影响产物与输入摘要",
    "随机种子",
    "收缩后的最小反例",
    "网页验证与隔离环境",
    "界面契约 · P0 关键旅程 · P1 回归测试",
    "问题驱动复审",
    "影响链分析",
    "实际差异复扫",
    "L1 / L2 / L3",
    "显性调用、隐性依赖",
    "风险适配，保护存量",
    "负向保证",
    "受保护历史测试",
    "调用边差异",
    "副作用预算",
    "新增测试证明新行为",
    "保护旧根基",
    "工作包上下文预算",
    "活动状态与证据索引",
    "二维风险模型",
    "变更风险 L1—L4",
    "纯展示 / 一般交互",
    "开发已验证",
    "合并已验证",
    "发布已就绪",
    "门禁依赖图与精确复用",
    "p50/p90",
    "测试去重不降低质量",
    "贯穿案例",
    "分组演练",
)

ABSTRACT_ENGLISH_HEADINGS = (
    "Capability escalation",
    "Candidate-bound final",
    "Canonical sources",
    "Token routing",
    "Brownfield governance",
    "Project anatomy",
    "CLI map",
    "Adoption path",
    "Common misconceptions",
)

INTEGRATED_SEQUENCES = (
    ("阶段层", "二维风险模型", "工作包上下文预算", "六步主线的输入与输出"),
    ("设计评审（design-review）", "影响链触发与豁免", "构建 Build"),
    ("网页验证与隔离环境", "测试去重不降低质量", "代码评审（code-review）"),
    ("提交、发布与证据", "发布与恢复演练", "三步启动"),
    ("分组演练与行动清单", "课程总结"),
)


def extract_slide_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )


def shape_geometry(slide) -> list[tuple[int, int, int, int, int]]:
    return [
        (int(shape.shape_type), shape.left, shape.top, shape.width, shape.height)
        for shape in slide.shapes
    ]


def picture_geometry(slide) -> list[tuple[int, int, int, int]]:
    return [
        (shape.left, shape.top, shape.width, shape.height)
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]


def inside_canvas(shape, prs: Presentation) -> bool:
    return (
        shape.left >= 0
        and shape.top >= 0
        and shape.left + shape.width <= prs.slide_width
        and shape.top + shape.height <= prs.slide_height
    )


def expected_strings(spec: dict[str, Any]) -> list[str]:
    principles = spec.get("principles", [])
    return [
        str(spec["title"]),
        *(str(item["text"]) for item in spec.get("updates", [])),
        *(str(item[field]) for item in principles for field in ("number", "title", "description")),
    ]


def verify(source_path: Path = DEFAULT_SOURCE, deck_path: Path = DEFAULT_DECK) -> dict[str, Any]:
    errors: list[str] = []
    source = yaml.safe_load(source_path.read_text(encoding="utf-8"))
    manifest = yaml.safe_load((ROOT / "framework-manifest.yaml").read_text(encoding="utf-8"))
    v1 = Presentation(V1)
    prs = Presentation(deck_path)
    slides = source.get("slides", [])
    expected_count = int(source.get("slide_count", len(slides)))
    template_count = int(source.get("template", {}).get("slides", 0))
    titles = [str(slide.get("title", "")) for slide in slides]

    actual_v1_sha = sha256(V1.read_bytes()).hexdigest()
    if source.get("source_version") != manifest.get("version"):
        errors.append("source_version does not match framework manifest")
    if actual_v1_sha != V1_SHA256 or source.get("template", {}).get("sha256") != V1_SHA256:
        errors.append("V1 template checksum changed or source guard is incorrect")
    if len(slides) != expected_count or len(v1.slides) != template_count or len(prs.slides) != expected_count:
        errors.append(
            "slide count mismatch: "
            f"declared={expected_count}, source={len(slides)}, "
            f"template={len(v1.slides)}/{template_count}, output={len(prs.slides)}"
        )
    if (prs.slide_width, prs.slide_height) != (v1.slide_width, v1.slide_height):
        errors.append("V2.4 canvas does not match the V1 template")
    if abs(prs.slide_width / prs.slide_height - 16 / 9) >= 0.0001:
        errors.append("deck is not 16:9 within V1 template tolerance")
    for sequence in INTEGRATED_SEQUENCES:
        try:
            positions = [titles.index(title) for title in sequence]
        except ValueError as exc:
            errors.append(f"integrated sequence is missing title: {exc}")
        else:
            if positions != sorted(positions):
                errors.append("topics are not integrated in order: " + " → ".join(sequence))
    if "概览完成，进入机制与案例" in titles:
        errors.append("contains obsolete overview/deep-dive split page")

    all_text: list[str] = []
    editable = 0
    geometry_mismatches = 0
    logo_issues = 0
    background_issues = 0
    new_out_of_bounds = 0
    for index, (spec, revised) in enumerate(zip(slides, prs.slides), 1):
        template_number = int(spec.get("template_slide", 0))
        if not 1 <= template_number <= template_count:
            errors.append(f"slide {index} has invalid template_slide: {template_number}")
            continue
        original = v1.slides[template_number - 1]
        if spec.get("number") != index:
            errors.append(f"source slide {index} has incorrect number")
        text = extract_slide_text(revised)
        original_text = extract_slide_text(original)
        all_text.append(text)
        original_has_page_number = any(
            re.fullmatch(r"\d{1,3}\s*/\s*\d{1,3}", shape.text.strip())
            for shape in original.shapes
            if getattr(shape, "has_text_frame", False)
        )
        page_number = f"{index:02d} / {expected_count}"
        if original_has_page_number and page_number not in text:
            errors.append(f"slide {index} missing dynamic page number {page_number}")
        for expected in expected_strings(spec):
            if expected not in text:
                errors.append(f"slide {index} missing source text: {expected[:32]}")

        controlled_geometry_change = bool(spec.get("principles")) or template_number == 17
        if not controlled_geometry_change and shape_geometry(revised) != shape_geometry(original):
            geometry_mismatches += 1
        if picture_geometry(revised) != picture_geometry(original) or len(picture_geometry(revised)) != 1:
            logo_issues += 1
        original_background = original.element.cSld.bg
        revised_background = revised.element.cSld.bg
        if (original_background is None) != (revised_background is None) or (
            original_background is not None
            and revised_background is not None
            and original_background.xml != revised_background.xml
        ):
            background_issues += 1
        for revised_shape in revised.shapes:
            if not inside_canvas(revised_shape, prs):
                new_out_of_bounds += 1
        for revised_shape in revised.shapes:
            if getattr(revised_shape, "has_text_frame", False) and revised_shape.text.strip():
                editable += 1

    joined = "\n".join(all_text)
    for topic in REQUIRED_TOPICS:
        if topic not in joined:
            errors.append(f"missing required topic: {topic}")
    for claim in source.get("forbidden_claims", []):
        if claim in joined:
            errors.append(f"contains forbidden legacy claim: {claim}")
    for heading in ABSTRACT_ENGLISH_HEADINGS:
        if heading in joined:
            errors.append(f"contains abstract English heading: {heading}")
    for explanation in ("二维风险模型", "冻结候选", "开发已验证", "被覆盖"):
        if explanation not in joined:
            errors.append(f"missing plain-language explanation: {explanation}")
    for phase in ("draft", "proposal", "design", "build", "verify", "retro", "release", "complete", "archived"):
        if phase not in joined:
            errors.append(f"missing state schema phase: {phase}")
    if geometry_mismatches:
        errors.append(f"{geometry_mismatches} slides drift from V1 shape geometry")
    if logo_issues:
        errors.append(f"{logo_issues} slides lost or moved the Measures logo")
    if background_issues:
        errors.append(f"{background_issues} slides lost or changed their V1 background")
    if new_out_of_bounds:
        errors.append(f"{new_out_of_bounds} shapes newly exceed the V1 canvas")
    minimum_editable = expected_count * 5
    if editable < minimum_editable:
        errors.append(
            f"expected at least {minimum_editable} editable text shapes, got {editable}"
        )

    for spec in slides:
        if len(spec["title"]) > 34:
            errors.append(f"slide {spec['number']} title exceeds budget")
        updates = spec.get("updates", [])
        if len(updates) > 20 or len({item["shape"] for item in updates}) != len(updates):
            errors.append(f"slide {spec['number']} has invalid update slots")
        for update in updates:
            text = str(update["text"])
            extended_line = (spec["template_slide"], update["slot"]) in {
                (13, "section-note"), (20, "footer")
            }
            line_limit = 100 if extended_line else 64
            if len(text) > 100 or any(len(line) > line_limit for line in text.splitlines()):
                errors.append(f"slide {spec['number']} text exceeds budget: {update['slot']}")
        principles = spec.get("principles", [])
        if principles and len(principles) != 6:
            errors.append(f"slide {spec['number']} must contain exactly six principles")
        for principle in principles:
            if any(len(str(principle[field])) > 32 for field in ("number", "title", "description")):
                errors.append(f"slide {spec['number']} principle text exceeds budget")

    return {
        "ok": not errors,
        "source_version": source.get("source_version"),
        "slides": len(prs.slides),
        "editable_text_shapes": editable,
        "geometry_mismatches": geometry_mismatches,
        "logo_issues": logo_issues,
        "background_issues": background_issues,
        "new_out_of_bounds_shapes": new_out_of_bounds,
        "errors": errors,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--deck", type=Path, default=DEFAULT_DECK)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    result = verify(args.source.resolve(), args.deck.resolve())
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    elif result["ok"]:
        print(
            f"valid V1-structured MASE v{result['source_version']} deck: "
            f"{result['slides']} slides, {result['editable_text_shapes']} editable text shapes, "
            f"{result['logo_issues']} logo issues"
        )
    else:
        print("invalid MASE training deck")
        for error in result["errors"]:
            print(f"- {error}")
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
