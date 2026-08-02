#!/usr/bin/env python3
"""Build the editable MASE v2.4 deck by updating the guarded V1 template."""

from __future__ import annotations

import argparse
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
from typing import Any

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "training/mase-framework/mase-training-v2.4.yaml"
DEFAULT_OUTPUT = ROOT / "training/mase-framework/MASE框架培训讲义V2.4.pptx"


def load_source(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    manifest = yaml.safe_load((ROOT / "framework-manifest.yaml").read_text(encoding="utf-8"))
    if data.get("source_version") != manifest.get("version"):
        raise ValueError(
            f"training source {data.get('source_version')} does not match "
            f"framework {manifest.get('version')}"
        )
    slides = data.get("slides", [])
    template_count = int(data.get("template", {}).get("slides", 0))
    expected_count = int(data.get("slide_count", len(slides)))
    if template_count != 37 or len(slides) != expected_count:
        raise ValueError(
            f"expected {expected_count} slides, got source={len(slides)} template={template_count}"
        )
    if [slide.get("number") for slide in slides] != list(range(1, expected_count + 1)):
        raise ValueError(f"slide numbers must be continuous from 1 to {expected_count}")
    valid_tracks = {"core", "deep-dive", "exercise"}
    default_track = str(data.get("default_track", "core"))
    if default_track not in valid_tracks:
        raise ValueError(f"invalid default track: {default_track}")
    for slide in slides:
        slide.setdefault("track", default_track)
        if slide["track"] not in valid_tracks:
            raise ValueError(f"slide {slide['number']} has invalid track: {slide['track']}")
        template_slide = int(slide.get("template_slide", 0))
        if not 1 <= template_slide <= template_count:
            raise ValueError(
                f"slide {slide['number']} requires template_slide within 1..{template_count}"
            )
    principle_slides = [slide for slide in slides if slide.get("principles")]
    principles = principle_slides[0].get("principles", []) if len(principle_slides) == 1 else []
    if len(principles) != 6 or any(
        set(item) != {"number", "title", "description"} for item in principles
    ):
        raise ValueError("the deck must define exactly one six-principle slide")
    return data


def template_path(data: dict[str, Any]) -> Path:
    relative = Path(str(data["template"]["path"]))
    path = (ROOT / relative).resolve()
    try:
        path.relative_to(ROOT)
    except ValueError as exc:
        raise ValueError("training template escapes repository root") from exc
    if not path.is_file():
        raise FileNotFoundError(path)
    actual = sha256(path.read_bytes()).hexdigest()
    expected = str(data["template"]["sha256"])
    if actual != expected:
        raise ValueError(f"V1 template checksum changed: expected {expected}, got {actual}")
    return path


def replace_text_preserving_style(
    shape,
    new_text: str,
    *,
    slide_number: int,
    font_sizes: list[float] | None = None,
) -> None:
    if not getattr(shape, "has_text_frame", False):
        raise ValueError(f"slide {slide_number} target shape is not editable text")
    paragraphs = list(shape.text_frame.paragraphs)
    values = str(new_text).split("\n")
    if len(values) != len(paragraphs):
        raise ValueError(
            f"slide {slide_number} paragraph mismatch for {shape.name}: "
            f"template={len(paragraphs)} source={len(values)}"
        )
    if font_sizes is not None and len(font_sizes) != len(paragraphs):
        raise ValueError(
            f"slide {slide_number} font-size mismatch for {shape.name}: "
            f"paragraphs={len(paragraphs)} font_sizes={len(font_sizes)}"
        )
    for index, (paragraph, value) in enumerate(zip(paragraphs, values)):
        runs = list(paragraph.runs)
        if runs:
            runs[0].text = value
            for run in runs[1:]:
                run.text = ""
        else:
            runs = [paragraph.add_run()]
            runs[0].text = value
        if font_sizes is not None:
            size = float(font_sizes[index])
            if size <= 0:
                raise ValueError(f"slide {slide_number} font size must be positive")
            for run in runs:
                run.font.size = Pt(size)


def apply_slide_updates(slide, spec: dict[str, Any]) -> None:
    seen: set[int] = set()
    for update in spec.get("updates", []):
        shape_index = int(update["shape"])
        if shape_index in seen:
            raise ValueError(f"slide {spec['number']} repeats shape {shape_index}")
        if not 0 <= shape_index < len(slide.shapes):
            raise ValueError(f"slide {spec['number']} has no shape {shape_index}")
        replace_text_preserving_style(
            slide.shapes[shape_index],
            str(update["text"]),
            slide_number=int(spec["number"]),
            font_sizes=update.get("font_sizes"),
        )
        seen.add(shape_index)
    slide_text = "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )
    if str(spec["title"]) not in slide_text:
        raise ValueError(f"slide {spec['number']} does not contain declared title: {spec['title']}")


def _clone_shape(slide, shape):
    clone = deepcopy(shape.element)
    slide.shapes._spTree.insert_element_before(clone, "p:extLst")
    return slide.shapes[-1]


def clone_slide(prs: Presentation, source_slide):
    """Clone an approved V1 layout while preserving editable shapes and logo geometry."""

    destination = prs.slides.add_slide(prs.slide_layouts[6])
    source_background = source_slide.element.cSld.bg
    if source_background is not None:
        destination.element.cSld.insert(0, deepcopy(source_background))
    for placeholder in list(destination.shapes):
        destination.shapes._spTree.remove(placeholder.element)
    for shape in source_slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture = destination.shapes.add_picture(
                BytesIO(shape.image.blob), shape.left, shape.top, shape.width, shape.height
            )
            picture.crop_left = shape.crop_left
            picture.crop_right = shape.crop_right
            picture.crop_top = shape.crop_top
            picture.crop_bottom = shape.crop_bottom
        else:
            clone = deepcopy(shape.element)
            destination.shapes._spTree.insert_element_before(clone, "p:extLst")
    return destination


def update_page_number(slide, number: int, total: int) -> None:
    pattern = re.compile(r"^\d{1,3}\s*/\s*\d{1,3}$")
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and pattern.fullmatch(shape.text.strip()):
            replace_text_preserving_style(
                shape, f"{number:02d} / {total}", slide_number=number
            )


def keep_shapes_inside_canvas(slide, prs: Presentation) -> None:
    """Move inherited shapes inside the canvas without changing their size or style."""

    for shape in slide.shapes:
        if shape.width > prs.slide_width or shape.height > prs.slide_height:
            raise ValueError(f"shape {shape.name} is larger than the slide canvas")
        shape.left = min(max(shape.left, 0), prs.slide_width - shape.width)
        shape.top = min(max(shape.top, 0), prs.slide_height - shape.height)


def _remove_slide(prs: Presentation, index: int) -> None:
    slide_id = prs.slides._sldIdLst[index]
    relationship_id = slide_id.rId
    prs.part.drop_rel(relationship_id)
    del prs.slides._sldIdLst[index]


def apply_six_principle_layout(slide, spec: dict[str, Any]) -> None:
    principles = spec.get("principles")
    if not principles:
        return
    if len(principles) != 6:
        raise ValueError(f"slide {spec['number']} requires exactly six principles")

    groups = [
        tuple(slide.shapes[index] for index in range(start, start + 4))
        for start in (2, 6, 10, 14)
    ]
    template_group = groups[0]
    for _ in range(2):
        groups.append(tuple(_clone_shape(slide, shape) for shape in template_group))

    area_left = 731520
    area_width = 10728655
    column_gap = 180000
    card_width = (area_width - 2 * column_gap) // 3
    card_height = 1750000
    row_tops = (1950000, 4050000)
    column_lefts = tuple(
        area_left + column * (card_width + column_gap) for column in range(3)
    )

    for index, (group, principle) in enumerate(zip(groups, principles)):
        background, number, title, description = group
        left = column_lefts[index % 3]
        top = row_tops[index // 3]
        background.left, background.top = left, top
        background.width, background.height = card_width, card_height
        number.left, number.top = left + 250000, top + 170000
        number.width, number.height = 600000, 360000
        title.left, title.top = left + 250000, top + 520000
        title.width, title.height = card_width - 500000, 440000
        description.left, description.top = left + 250000, top + 1030000
        description.width, description.height = card_width - 500000, 500000
        replace_text_preserving_style(
            number, str(principle["number"]), slide_number=int(spec["number"])
        )
        replace_text_preserving_style(
            title, str(principle["title"]), slide_number=int(spec["number"])
        )
        replace_text_preserving_style(
            description,
            str(principle["description"]),
            slide_number=int(spec["number"]),
        )


def build_deck(source: Path, output: Path, *, track: str = "full") -> Path:
    data = load_source(source)
    template = template_path(data)
    prs = Presentation(template)
    template_count = int(data["template"]["slides"])
    if len(prs.slides) != template_count:
        raise ValueError(f"V1 template must contain 37 slides, got {len(prs.slides)}")

    template_slides = list(prs.slides)[:template_count]
    selected_specs = list(data["slides"])
    if track != "full":
        if track not in {"core", "deep-dive", "exercise"}:
            raise ValueError(f"unsupported training track: {track}")
        selected_specs = [
            spec for spec in data["slides"]
            if spec["track"] == "core" or spec["track"] == track
        ]

    for spec in selected_specs:
        slide = clone_slide(prs, template_slides[int(spec["template_slide"]) - 1])
        apply_slide_updates(slide, spec)
        apply_six_principle_layout(slide, spec)

    for index in range(template_count - 1, -1, -1):
        _remove_slide(prs, index)

    total = len(prs.slides)
    for number, slide in enumerate(prs.slides, 1):
        update_page_number(slide, number, total)
        keep_shapes_inside_canvas(slide, prs)

    prs.core_properties.title = data["title"]
    prs.core_properties.subject = "MASE v2.4 training in the V1 teaching structure"
    prs.core_properties.author = "Measures · MASE"
    prs.core_properties.keywords = "MASE,v2.4,V1-structure,risk-adaptive,TDD,contract,evidence"
    prs.core_properties.comments = (
        "Generated from the checksum-guarded V1 template and mase-training-v2.4.yaml; "
        "canonical rules remain in runtime sources."
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--track", choices=("full", "core", "deep-dive", "exercise"), default="full")
    args = parser.parse_args()
    path = build_deck(args.source.resolve(), args.output.resolve(), track=args.track)
    print(f"built {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
