from hashlib import sha256
from pathlib import Path
import re

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from scripts.build_mase_training_deck import build_deck


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "training/mase-framework/mase-training-v2.4.yaml"
OUTPUT = ROOT / "training/mase-framework/MASE框架培训讲义V2.4.pptx"
V1 = ROOT / "training/mase-framework/MASE框架培训讲义V1.pptx"
V1_SHA256 = "fcdf4bd0b8fe3adb5316c8dbccf6e70d91b0274857a5f97827709053886fe188"

REQUIRED_TOPICS = (
    "MASE v2.4",
    "风险自适应",
    "Lite",
    "Standard",
    "Strict",
    "六步开发主线",
    "三种过程档位",
    "轻量档（Lite）",
    "标准档（Standard）",
    "严格档（Strict）",
    "Capability",
    "门禁运行器",
    "候选冻结",
    "mase-state.yaml",
    ".mase/gates.yaml",
    "Brownfield",
    "隔离环境",
    "上下文路由",
    "属性测试",
    "最小反例",
    "可选发布附加流程",
    "打包、发布、部署、线上验证或恢复",
    "不可变制品",
    "target_ready",
    "live_verified",
    "release-software",
    "观察与恢复",
    "六项开发原则",
    "验证职责分层",
    "原则落地：三次回看",
    "Spec 与测试设计",
    "独立测试用例用 Test ID 追踪",
    "风险与契约驱动设计",
    "证据新鲜度闭环",
    "需求到证据的追踪链",
    "候选 → 制品 → 目标 → 线上 → 观察",
    "门禁计划",
    "固定代码、影响产物与输入摘要",
    "随机种子",
    "收缩后的最小反例",
    "网页验证与隔离环境",
    "界面契约 · P0 关键旅程 · P1 回归测试",
    "问题驱动复审",
    "影响链分析",
    "实际差异复扫",
    "L1 / L2 / L3",
    "显性调用、隐性依赖",
    "风险适配，保护存量",
    "负向保证",
    "受保护历史测试",
    "调用边差异",
    "副作用预算",
    "新增测试证明新行为",
    "保护旧根基",
    "工作包上下文预算",
    "活动状态与证据索引",
    "测试去重不降低质量",
    "二维风险模型",
    "变更风险 L1—L4",
    "纯展示 / 一般交互",
    "开发已验证",
    "合并已验证",
    "发布已就绪",
    "门禁依赖图与精确复用",
    "p50/p90",
    "让 Agentic Coding 高效交付正确、健壮、优化且易于维护的代码",
    "高效交付 · 需求正确 · 运行健壮 · 质量优化 · 整洁可维护",
    "本仓样本约 71%",
    "1 - 17.2 / 58.4",
    "贯穿案例",
    "分组演练",
)

FORBIDDEN_CLAIMS = (
    "每 20 次对话",
    "人不读需求文档",
    "测试绿灯 = 实现正确",
    "GWT 用例 / 整个流水线的唯一真相源",
    "六阶段 · 层层递进",
    "cp -r openspec/changes/_template/",
    "六步教学主线",
    "三 Profile",
    "Release Overlay",
    "Web 验证与 Sandbox",
    "UI contract",
    "P0 journey",
    "P1 regression",
    "保存 seed",
    "shrink",
    "fixture",
    "adapter",
    "flaky",
    "trace/",
    "AI 写的测试绝不可信",
)

ABSTRACT_ENGLISH_HEADINGS = (
    "Capability escalation",
    "Candidate-bound final",
    "Canonical sources",
    "Token routing",
    "Brownfield governance",
    "Project anatomy",
    "CLI map",
    "Adoption path",
    "Common misconceptions",
)


def deck_text(prs: Presentation) -> str:
    return "\n".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )


def slide_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )


def slide_for_title(prs: Presentation, title: str):
    source = yaml.safe_load(SOURCE.read_text(encoding="utf-8"))
    index = next(index for index, slide in enumerate(source["slides"]) if slide["title"] == title)
    return prs.slides[index]


def shape_geometry(slide) -> list[tuple[int, int, int, int, int]]:
    return [
        (int(shape.shape_type), shape.left, shape.top, shape.width, shape.height)
        for shape in slide.shapes
    ]


def picture_geometry(slide) -> list[tuple[int, int, int, int]]:
    return [
        (shape.left, shape.top, shape.width, shape.height)
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]


def test_v1_training_deck_is_preserved_byte_for_byte():
    assert sha256(V1.read_bytes()).hexdigest() == V1_SHA256


def test_v23_source_and_editable_deck_exist_and_match_manifest_version():
    source = yaml.safe_load(SOURCE.read_text(encoding="utf-8"))
    manifest = yaml.safe_load((ROOT / "framework-manifest.yaml").read_text())
    prs = Presentation(OUTPUT)
    v1 = Presentation(V1)

    assert source["source_version"] == manifest["version"] == "2.4.0"
    assert source["slide_count"] == len(source["slides"]) == len(prs.slides) == 66
    assert source["template"]["slides"] == len(v1.slides) == 37
    assert (prs.slide_width, prs.slide_height) == (v1.slide_width, v1.slide_height)
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.0001
    assert all(slide["number"] == index for index, slide in enumerate(source["slides"], 1))
    assert all(slide.get("track", source["default_track"]) in {"core", "deep-dive", "exercise"} for slide in source["slides"])
    assert all(1 <= slide["template_slide"] <= 37 for slide in source["slides"])


def test_v23_deck_covers_current_topics_and_rejects_legacy_claims():
    text = deck_text(Presentation(OUTPUT))

    for topic in REQUIRED_TOPICS:
        assert topic in text
    for claim in FORBIDDEN_CLAIMS:
        assert claim not in text
    assert "产品有 UI 且本 change 修改 UI" in text
    assert "补充而不替代" in text
    assert "六步开发主线（不是状态枚举）" in text
    for phase in ("draft", "proposal", "design", "build", "verify", "retro", "release", "complete", "archived"):
        assert phase in text


def test_v23_deck_is_numbered_editable_and_inside_canvas():
    prs = Presentation(OUTPUT)
    v1 = Presentation(V1)
    source = yaml.safe_load(SOURCE.read_text(encoding="utf-8"))
    editable_text_shapes = 0

    for index, (spec, slide) in enumerate(zip(source["slides"], prs.slides), 1):
        template_number = spec["template_slide"]
        original = v1.slides[template_number - 1]
        original_has_page_number = any(
            re.fullmatch(r"\d{1,3}\s*/\s*\d{1,3}", shape.text.strip())
            for shape in original.shapes
            if getattr(shape, "has_text_frame", False)
        )
        page_number = f"{index:02d} / 66"
        if original_has_page_number:
            assert page_number in slide_text(slide)
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                editable_text_shapes += 1

    assert editable_text_shapes >= 330


def test_core_principles_are_front_loaded_and_later_slides_are_mechanisms():
    prs = Presentation(OUTPUT)
    principles = slide_for_title(prs, "六项开发原则")

    assert "核心理念" in slide_text(principles)
    assert "六项开发原则" in slide_text(principles)
    for number in range(1, 7):
        assert f"{number:02d}" in slide_text(principles)
    evidence = slide_for_title(prs, "证据新鲜度闭环")
    traceability = slide_for_title(prs, "需求到证据的追踪链")
    assert "核心理念" not in slide_text(evidence)
    assert "证据机制" in slide_text(evidence)
    assert "核心理念" not in slide_text(traceability)
    assert "追踪机制" in slide_text(traceability)


def test_evidence_circles_and_review_guidance_use_meaningful_chinese():
    prs = Presentation(OUTPUT)
    evidence = slide_for_title(prs, "证据新鲜度闭环")

    assert [evidence.shapes[index].text for index in (3, 7, 11, 15)] == [
        "输入",
        "执行",
        "证据",
        "新鲜度",
    ]
    for index in (3, 7, 11, 15):
        runs = evidence.shapes[index].text_frame.paragraphs[0].runs
        assert all(run.font.size and run.font.size.pt == 24 for run in runs)
    slide_30 = slide_text(slide_for_title(prs, "分级 TDD 与属性测试"))
    assert "随机种子" in slide_30 and "收缩后的最小反例" in slide_30
    assert "seed" not in slide_30 and "shrink" not in slide_30
    slide_31 = slide_text(slide_for_title(prs, "网页验证与隔离环境"))
    for english in ("Web", "Sandbox", "contract", "journey", "regression", "fixture", "adapter", "flaky", "trace"):
        assert english not in slide_31
    assert "构建阶段 · 技能" in slide_31
    assert "技能 · 网页测试" in slide_31
    slide_32 = slide_text(slide_for_title(prs, "代码评审（code-review）"))
    assert "问题驱动复审" in slide_32
    assert "首轮无异议即结束" in slide_32


def test_v23_preserves_v1_page_slots_shape_geometry_and_measures_logo():
    v1 = Presentation(V1)
    v23 = Presentation(OUTPUT)
    source = yaml.safe_load(SOURCE.read_text(encoding="utf-8"))

    assert (v23.slide_width, v23.slide_height) == (v1.slide_width, v1.slide_height)
    for index, (spec, revised) in enumerate(zip(source["slides"], v23.slides), 1):
        template_number = spec["template_slide"]
        original = v1.slides[template_number - 1]
        if spec.get("principles"):
            assert shape_geometry(revised) != shape_geometry(original)
            assert len(revised.shapes) == len(original.shapes) + 8
        elif template_number == 17:
            assert shape_geometry(revised) != shape_geometry(original)
            assert revised.shapes[14].top + revised.shapes[14].height <= v23.slide_height
        else:
            assert shape_geometry(revised) == shape_geometry(original), f"slide {index} geometry drift"
        assert picture_geometry(revised) == picture_geometry(original), f"slide {index} logo drift"
        assert len(picture_geometry(revised)) == 1, f"slide {index} Measures logo missing"
        assert revised.element.cSld.bg.xml == original.element.cSld.bg.xml, (
            f"slide {index} background drift"
        )

    for expected in ("Agent 层", "阶段层", "Skills 矩阵", "快速上手"):
        assert any(expected in slide_text(slide) for slide in v23.slides)


def test_v23_uses_chinese_first_headings_and_concrete_explanations():
    text = deck_text(Presentation(OUTPUT))

    for heading in ABSTRACT_ENGLISH_HEADINGS:
        assert heading not in text
    for explanation in (
        "二维风险模型",
        "冻结候选",
        "开发已验证",
        "被覆盖",
    ):
        assert explanation in text


def test_v23_source_respects_content_budgets():
    source = yaml.safe_load(SOURCE.read_text(encoding="utf-8"))

    assert source["template"]["sha256"] == V1_SHA256
    assert source["template"]["slides"] == 37
    for slide in source["slides"]:
        assert len(slide["title"]) <= 34
        updates = slide.get("updates", [])
        assert len(updates) <= 20
        assert len({update["shape"] for update in updates}) == len(updates)
        assert all(update["slot"] and update["text"] for update in updates)
        for update in updates:
            assert len(update["text"]) <= 100
            extended_line = (slide["template_slide"], update["slot"]) in {
                (13, "section-note"), (20, "footer")
            }
            line_limit = 100 if extended_line else 64
            assert all(len(line) <= line_limit for line in update["text"].splitlines())
        principles = slide.get("principles", [])
        if principles:
            assert len(principles) == 6
            assert [item["number"] for item in principles] == [f"{number:02d}" for number in range(1, 7)]
            assert all(len(item["title"]) <= 16 for item in principles)
            assert all(len(item["description"]) <= 32 for item in principles)


def test_v23_dense_vv_panels_use_smaller_detail_text():
    prs = Presentation(OUTPUT)

    slide_16_detail = slide_for_title(prs, "按风险做设计").shapes[23].text_frame.paragraphs[1]
    assert all(run.font.size and run.font.size.pt == 13 for run in slide_16_detail.runs)

    slide_17_panel = slide_for_title(prs, "构建 Build").shapes[13].text_frame
    assert slide_17_panel.paragraphs[0].text == "V&V：新增测试证明新行为；负向保证保护旧根基"
    assert all(run.font.size and run.font.size.pt == 18 for run in slide_17_panel.paragraphs[0].runs)
    assert slide_17_panel.paragraphs[1].text == "仍要对照 Spec、调用边、受保护测试和副作用预算"


def test_training_tracks_are_generated_from_the_same_source(tmp_path):
    expected_counts = {"core": 37, "deep-dive": 65, "exercise": 38}

    for track, expected_count in expected_counts.items():
        output = tmp_path / f"mase-{track}.pptx"
        build_deck(SOURCE, output, track=track)
        prs = Presentation(output)

        assert len(prs.slides) == expected_count
        assert any(
            re.search(rf"\b\d{{2}} / {expected_count}\b", slide_text(slide))
            for slide in prs.slides
        )


def test_deep_dive_pages_are_integrated_with_their_chapter_anchors():
    source = yaml.safe_load(SOURCE.read_text(encoding="utf-8"))
    titles = [slide["title"] for slide in source["slides"]]

    def position(title: str) -> int:
        return titles.index(title)

    assert position("二维风险模型") < position("需求 Proposal")
    assert position("工作包上下文预算") < position("六步主线的输入与输出")
    assert position("设计评审（design-review）") < position("影响链触发与豁免") < position("构建 Build")
    assert position("网页验证与隔离环境") < position("测试去重不降低质量") < position("代码评审（code-review）")
    assert position("提交、发布与证据") < position("发布与恢复演练") < position("三步启动")
    assert position("分组演练与行动清单") < position("课程总结")
    assert "概览完成，进入机制与案例" not in titles
