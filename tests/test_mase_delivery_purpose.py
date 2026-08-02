from pathlib import Path

from pptx import Presentation

from scripts.estimate_mase_speed_gain import estimate_from_repository


ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "training/mase-framework/MASE框架培训讲义V2.4.pptx"
PURPOSE = "让 Agentic Coding 高效交付正确、健壮、优化且易于维护的代码"
OUTCOMES = ("高效交付", "需求正确", "运行健壮", "质量优化", "整洁可维护")


def _deck_text() -> str:
    return "\n".join(
        shape.text
        for slide in Presentation(DECK).slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def test_design_purpose_is_consistent_across_canonical_surfaces():
    paths = (
        "project-rules.md",
        "README.md",
        "docs/MASE-framework.md",
        "docs/design-principles.md",
        "docs/user-guide.md",
    )

    for relative in paths:
        text = (ROOT / relative).read_text(encoding="utf-8")
        assert PURPOSE in text, relative
        for outcome in OUTCOMES:
            assert outcome in text, f"{relative}: {outcome}"

    deck = _deck_text()
    assert PURPOSE in deck
    for outcome in OUTCOMES:
        assert outcome in deck

    for relative in ("AGENTS.md", "CLAUDE.md", "CONVENTIONS.md", ".github/copilot-instructions.md"):
        adapter = (ROOT / relative).read_text(encoding="utf-8")
        assert PURPOSE in adapter, relative
        for outcome in OUTCOMES:
            assert outcome in adapter, f"{relative}: {outcome}"


def test_repository_speed_estimate_is_reproducible_and_bounded():
    report = estimate_from_repository(ROOT)

    assert report["automatic_evidence_records"] >= 54
    assert 15 <= report["development_observed_median_serial_seconds"] <= 20
    assert 50 <= report["old_pre_hand_test_observed_median_serial_seconds"] <= 70
    assert 0.65 <= report["hand_test_wait_reduction_estimate"] <= 0.75
    assert report["percentile_ready"] is False
    assert report["confidence"] == "low"
    assert len(report["insufficient_sample_gates"]) == 5
    assert report["universal_commitment"] is False
